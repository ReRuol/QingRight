"""生成最小合法的 blank.docx / blank.xlsx（Word/Excel/Pages/Numbers 都能打开）。"""
import sys, zipfile
out = sys.argv[1]
CT = '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
def write(name, files):
    with zipfile.ZipFile(f"{out}/{name}", "w", zipfile.ZIP_DEFLATED) as z:
        for k, v in files.items(): z.writestr(k, CT + v)
rels = lambda t, tgt: f'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="{t}" Target="{tgt}"/></Relationships>'
OD = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument"
write("blank.docx", {
    "[Content_Types].xml": '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>',
    "_rels/.rels": rels(OD, "word/document.xml"),
    "word/document.xml": '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p/></w:body></w:document>',
})
write("blank.xlsx", {
    "[Content_Types].xml": '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/><Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/></Types>',
    "_rels/.rels": rels(OD, "xl/workbook.xml"),
    "xl/workbook.xml": '<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><sheets><sheet name="Sheet1" sheetId="1" r:id="rId1"/></sheets></workbook>',
    "xl/_rels/workbook.xml.rels": rels("http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet", "worksheets/sheet1.xml"),
    "xl/worksheets/sheet1.xml": '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData/></worksheet>',
})

# 纯文本类空白模板
open(f"{out}/blank.rtf", "w").write("{\\rtf1\\ansi\\ansicpg936\\deff0{\\fonttbl{\\f0 PingFangSC-Regular;}}\\f0\\fs24\n}\n")

# 一页空白 A4 PDF，手写交叉引用表，偏移按字节算
objs = [b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Contents 4 0 R /Resources << >> >>",
        b"<< /Length 0 >>\nstream\n\nendstream"]
pdf = b"%PDF-1.4\n"; offs = []
for i, o in enumerate(objs, 1):
    offs.append(len(pdf)); pdf += f"{i} 0 obj\n".encode() + o + b"\nendobj\n"
x = len(pdf)
pdf += f"xref\n0 {len(objs)+1}\n0000000000 65535 f \n".encode() + b"".join(f"{o:010d} 00000 n \n".encode() for o in offs)
pdf += f"trailer\n<< /Size {len(objs)+1} /Root 1 0 R >>\nstartxref\n{x}\n%%EOF\n".encode()
open(f"{out}/blank.pdf", "wb").write(pdf)

# pptx 结构复杂（母版/版式/主题），交给 python-pptx；没装就跳过，保留仓库里已有的 blank.pptx
try:
    from pptx import Presentation
    from pptx.util import Inches
    p = Presentation(); p.slide_width, p.slide_height = Inches(13.333), Inches(7.5)  # 16:9
    p.slides.add_slide(p.slide_layouts[6])
    p.save(f"{out}/blank.pptx")
except ImportError:
    print("未安装 python-pptx，跳过 blank.pptx")
